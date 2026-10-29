#!/usr/bin/env python
"""export_papers.py 离线 smoke test —— **不联网、不碰真实 data/ 与 output/**。

覆盖（父任务验收清单 ①–⑦）:
  ① 只导出 approved（pending_review / rejected 被排除；`--include-pending` 才纳入）
  ② docx 生成且非空、能被 python-docx 重新打开并含题号
  ③ xlsx 行数 == 题数（+1 表头）
  ④ csv 以 UTF-8 BOM（\\xef\\xbb\\xbf）开头
  ⑤ PDF 走中文字体分支（字体缺失分支也断言：给警告 + 不生成 PDF）
  ⑥ genanki 生成 .apkg（存在且 > 0 字节）
  ⑦ `--dry-run` 不写任何文件

⚠️ 不使用 tempfile.TemporaryDirectory()：本机沙箱拒绝对系统临时目录做删除（WinError 5），
   临时文件统一放在工作区内 tests/.tmp/export_papers/ 下。
⚠️ tests/.tmp 是多个测试文件共享的目录：本文件只建/只清**自己的子目录** export_papers/。

运行: python tests/test_export_papers.py
"""

from __future__ import annotations

import io
import json
import shutil
import sqlite3
import subprocess
import sys
import zipfile
from contextlib import redirect_stdout
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

import export_papers as ep  # noqa: E402
import init_db  # noqa: E402  —— 复用 scripts/init_db.py 的 questions schema

PY = sys.executable
TMPROOT = ROOT / "tests" / ".tmp" / "export_papers"
E2E_DATE = "20260101"


# --------------------------------------------------------------------- 工具

def fresh_dir(name: str) -> Path:
    """在工作区内创建干净的临时目录（只清自己的子目录）。"""
    d = TMPROOT / name
    shutil.rmtree(d, ignore_errors=True)
    d.mkdir(parents=True, exist_ok=True)
    return d


def cleanup(d: Path) -> None:
    shutil.rmtree(d, ignore_errors=True)


def q(uid: str, qtype: str, stem: str, *, subject: str, status: str = "approved",
      options=None, answer=None, explanation=None, difficulty="中",
      code_template=None, knowledge_point=None) -> dict:
    """构造一行 questions（只填导出关心的列）。"""
    return {
        "question_uid": uid, "subject": subject, "question_type": qtype, "stem": stem,
        "status": status,
        "options": json.dumps(options, ensure_ascii=False) if options is not None else None,
        "answer": answer, "explanation": explanation, "difficulty": difficulty,
        "code_template": code_template, "knowledge_point": knowledge_point,
        "source_type": "generate",
        "source_url": f"https://example.invalid/{uid}",   # 明显的非真实域名（测试数据）
        "fact_source": "AI生成-未验证", "fact_verified": 0,
    }


def make_db(path: Path, rows: list[dict]) -> Path:
    """用 scripts/init_db.py 的 schema 造一个临时科目库。"""
    path.parent.mkdir(parents=True, exist_ok=True)
    conn = sqlite3.connect(str(path))
    conn.executescript(init_db.SUBJECT_SCHEMA)
    for row in rows:
        cols = ",".join(row)
        ph = ",".join("?" * len(row))
        conn.execute(f"INSERT INTO questions({cols}) VALUES({ph})", list(row.values()))
    conn.commit()
    conn.close()
    return path


def run_cli(*args: str) -> subprocess.CompletedProcess:
    return subprocess.run(
        [PY, str(ROOT / "scripts" / "export_papers.py"), *args],
        capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=300, cwd=str(ROOT),
    )


def last_json(stdout: str) -> dict:
    for line in reversed([ln for ln in stdout.splitlines() if ln.strip()]):
        try:
            obj = json.loads(line)
        except ValueError:
            continue
        if isinstance(obj, dict) and "papers" in obj:
            return obj
    raise AssertionError("stdout 里找不到 JSON 汇总行")


# --------------------------------------------------------------------- 固件

PENDING_STEM = "PENDING-EN-1 语法填空占位题干"
REJECTED_STEM = "REJECTED-EN-1 语法填空占位题干"
ACA_PENDING_STEM = "ACA-PENDING-1 单选占位题干"


def build_all_fixtures(data_dir: Path) -> None:
    """english / c_wuhan / aca_llm 三库；c_hbpu 故意缺失（覆盖「库不存在」分支）。

    配额截取场景：english.writing 的 paper_count=1，但给 2 道 approved → 应截掉 1 道。
    """
    make_db(data_dir / "english.sqlite", [
        q("en-g1", "grammar_fill", "EN-APPROVED-G1 语法填空题干（the ____ of）", subject="english"),
        q("en-g2", "grammar_fill", "EN-APPROVED-G2 语法填空题干（a ____ book）", subject="english"),
        q("en-g3", "grammar_fill", "EN-APPROVED-G3 语法填空题干（if I ____）", subject="english"),
        q("en-r1", "reading_fill", "EN-APPROVED-R1 阅读填空题干（Passage One）", subject="english",
          options={"1": "because", "2": "however", "3": "which"}),
        q("en-s1", "sentence_build", "EN-APPROVED-S1 连词成句：apple / an / day / keeps", subject="english"),
        q("en-w1", "writing", "EN-APPROVED-W1 写作：Write about your campus life.", subject="english", answer="范文一"),
        q("en-w2", "writing", "EN-APPROVED-W2 写作：Write about a person you admire.", subject="english", answer="范文二"),
        q("en-p1", "grammar_fill", PENDING_STEM, subject="english", status="pending_review"),
        q("en-x1", "grammar_fill", REJECTED_STEM, subject="english", status="rejected"),
    ])
    make_db(data_dir / "c_wuhan.sqlite", [
        q("cw-c1", "code_correction", "CW-APPROVED-C1 改错：找出下面程序中的错误。", subject="c_wuhan",
          code_template="#include <stdio.h>\nint main(){\n    int a = 1\n    printf(\"%d\", a);\n    return 0;\n}",
          knowledge_point="语法基础"),
        q("cw-p1", "programming", "CW-APPROVED-P1 编程：求 1 到 n 的和。", subject="c_wuhan",
          code_template="#include <stdio.h>\nint main(){\n    /* TODO */\n    return 0;\n}",
          answer="#include <stdio.h>\nint main(){ int n,s=0; scanf(\"%d\",&n); for(int i=1;i<=n;i++) s+=i; printf(\"%d\",s); return 0; }",
          explanation="循环累加即可。"),
    ])
    make_db(data_dir / "aca_llm.sqlite", [
        q("aca-s1", "single_choice", "ACA-S1 大模型概念：以下哪项属于提示词工程？", subject="aca_llm",
          options={"A": "温度参数", "B": "Few-shot 示例", "C": "学习率", "D": "批大小"},
          answer="B", explanation="Few-shot 属于提示词工程手段。"),
        q("aca-s2", "single_choice", "ACA-S2 微调：LoRA 的主要作用是？", subject="aca_llm",
          options={"A": "降低显存占用", "B": "提高温度", "C": "增加上下文", "D": "替换分词器"},
          answer="A", explanation="低秩适配减少可训练参数。", difficulty="易"),
        q("aca-m1", "multiple_choice", "ACA-M1 安全合规：以下哪些属于数据合规要求？", subject="aca_llm",
          options={"A": "最小够用原则", "B": "脱敏处理", "C": "随意共享", "D": "留存审计"},
          answer="ABD", explanation="C 明显违规。", difficulty="难"),
        q("aca-p1", "single_choice", ACA_PENDING_STEM, subject="aca_llm", status="pending_review",
          options={"A": "甲", "B": "乙"}),
    ])


_E2E: dict = {}


def e2e() -> dict:
    """跑一次完整的 CLI（all 科目）到临时 outdir，供多个测试共用（懒加载）。"""
    if _E2E:
        return _E2E
    data_dir = fresh_dir("e2e_data")
    out_base = fresh_dir("e2e_out")
    build_all_fixtures(data_dir)
    proc = run_cli("--data-dir", str(data_dir), "--outdir", str(out_base), "--date", E2E_DATE)
    assert proc.returncode == 0, f"CLI 退出码 {proc.returncode}\nSTDERR:\n{proc.stderr[-1200:]}"
    summary = last_json(proc.stdout)
    _E2E.update({"proc": proc, "summary": summary, "data_dir": data_dir,
                 "out": out_base / E2E_DATE, "out_base": out_base})
    return _E2E


# --------------------------------------------------------------------- ① 只导 approved

def test_fetch_questions_only_approved() -> None:
    d = fresh_dir("approved_only")
    try:
        db = make_db(d / "english.sqlite", [
            q("a1", "grammar_fill", "APPROVED-1", subject="english"),
            q("a2", "grammar_fill", "APPROVED-2", subject="english"),
            q("p1", "grammar_fill", "PENDING-1", subject="english", status="pending_review"),
            q("x1", "grammar_fill", "REJECTED-1", subject="english", status="rejected"),
            q("r1", "grammar_fill", "RAW-1", subject="english", status="raw"),
        ])
        rows = ep.fetch_questions(db, [ep.STATUS_APPROVED])
        assert len(rows) == 2, f"只应取 approved 2 题，实际 {len(rows)}"
        assert {r["question_uid"] for r in rows} == {"a1", "a2"}
        assert all(r["status"] == "approved" for r in rows)
        assert ep.statuses_for(False) == ["approved"]
        rows2 = ep.fetch_questions(db, ep.statuses_for(True))
        assert len(rows2) == 3, f"--include-pending 应取 3 题，实际 {len(rows2)}"
    finally:
        cleanup(d)


def test_pending_stem_absent_from_outputs() -> None:
    """端到端：pending / rejected 的题干**不得**出现在任何导出文件里。"""
    e = e2e()
    csv_text = (e["out"] / "all_questions.csv").read_text(encoding="utf-8-sig")
    assert "EN-APPROVED-G1" in csv_text, "approved 题应出现在 csv 里"
    for bad in (PENDING_STEM, REJECTED_STEM, ACA_PENDING_STEM):
        assert bad not in csv_text, f"非 approved 题不该出现在 csv：{bad}"
    assert e["summary"]["papers"]["english"]["statuses"] == ["approved"]


# --------------------------------------------------------------------- ② DOCX

def test_docx_generated_and_reopenable() -> None:
    from docx import Document

    e = e2e()
    for subject in ("english", "c_wuhan", "aca_llm"):
        path = e["out"] / f"{subject}_paper.docx"
        assert path.is_file(), f"缺 {path.name}"
        assert path.stat().st_size > 0, f"{path.name} 是空文件"
        doc = Document(str(path))
        text = "\n".join(p.text for p in doc.paragraphs)
        assert len(text.strip()) > 50, f"{path.name} 内容过少"
        assert "1." in text, f"{path.name} 缺题号"
        assert "试卷" in text, f"{path.name} 缺试卷头"
    en_text = "\n".join(p.text for p in Document(str(e["out"] / "english_paper.docx")).paragraphs)
    assert "EN-APPROVED-G1" in en_text, "docx 应含 approved 题干"
    assert "实收 6 题 / 目标 43 题" in en_text, f"试卷头应写实收/目标，实际片段：{[l for l in en_text.splitlines() if '实收' in l]}"
    assert PENDING_STEM not in en_text, "pending 题不得进试卷"


def test_paper_title_and_sections_from_taxonomy() -> None:
    e = e2e()
    ata = ep.load_taxonomy()
    doc_model = ep.build_document(
        "english", ep.subject_meta(ata, "english"), ep.question_types(ata, "english"),
        ep.fetch_questions(e["data_dir"] / "english.sqlite", ["approved"]),
        statuses=["approved"], generated_at="2026-01-01 00:00", warnings=[],
    )
    stats = doc_model["stats"]
    by_type = stats["by_type"]
    assert list(by_type) == ["grammar_fill", "sentence_build", "reading_fill", "reading_qa", "translation_fill", "writing"], \
        f"分节顺序必须与 taxonomy 的 question_types 顺序一致，实际 {list(by_type)}"
    assert by_type["grammar_fill"]["count"] == 3
    assert by_type["writing"]["target"] == 1 and by_type["writing"]["count"] == 1, "writing 应按 paper_count=1 截取"
    assert by_type["writing"]["dropped_by_quota"] == 1
    assert by_type["reading_qa"]["count"] == 0, "库里没有的题型应留空节"
    assert stats["actual_items"] == 6 and stats["target_items"] == 43
    section_texts = [b["text"] for b in doc_model["blocks"] if b["kind"] == "section"]
    assert any("语法填空" in t for t in section_texts) and any("写作" in t for t in section_texts)
    assert all("未获取" not in t for t in section_texts), "英语分值已知，不该出现未获取"


# --------------------------------------------------------------------- ③ XLSX

def test_xlsx_row_count_equals_question_count() -> None:
    from openpyxl import load_workbook

    e = e2e()
    path = e["out"] / "all_questions.xlsx"
    assert path.is_file() and path.stat().st_size > 0
    ws = load_workbook(str(path)).active
    expected = sum(p["actual_items"] for p in e["summary"]["papers"].values())
    assert ws.max_row - 1 == expected, f"xlsx 题行数应为 {expected}，实际 {ws.max_row - 1}"
    assert [c.value for c in ws[1]] == ep.EXPORT_HEADERS
    uid_col = [ws.cell(row=r, column=1).value for r in range(2, ws.max_row + 1)]
    assert len(set(uid_col)) == len(uid_col), "xlsx 里 uid 不应重复"


# --------------------------------------------------------------------- ④ CSV BOM

def test_csv_has_utf8_bom() -> None:
    e = e2e()
    path = e["out"] / "all_questions.csv"
    raw = path.read_bytes()
    assert raw[:3] == b"\xef\xbb\xbf", f"csv 必须以 UTF-8 BOM 开头，实际 {raw[:3]!r}"
    assert path.stat().st_size > 3
    header = raw.decode("utf-8-sig").splitlines()[0]
    assert header.split(",")[0] == "uid", f"表头异常：{header}"


# --------------------------------------------------------------------- ⑤ PDF / 中文字体

def test_pdf_cjk_font_branch() -> None:
    font, bold = ep.find_cjk_font()
    e = e2e()
    if not font:
        # 字体缺失分支：不生成 PDF + 明确警告（跳过项里必须有原因）
        assert not list(e["out"].glob("*.pdf")), "无中文字体时不得生成 PDF"
        reasons = [s["reason"] for s in e["summary"]["skipped"] if s["item"].endswith(".pdf")]
        assert reasons and any("字体" in r for r in reasons), f"缺字体必须有明确警告，实际 {reasons}"
        assert any("字体" in w for w in e["summary"]["warnings"])
        return
    # 字体存在分支：PDF 必须生成、非空、且中文**可被提取**（不是方块）
    for subject in ("english", "c_wuhan", "aca_llm"):
        path = e["out"] / f"{subject}_paper.pdf"
        assert path.is_file(), f"缺 {path.name}（字体 {font} 可用时应当生成）"
        assert path.stat().st_size > 0, f"{path.name} 是空文件"
        assert path.read_bytes()[:5] == b"%PDF-", f"{path.name} 不是 PDF"
    try:
        import pdfplumber
    except ImportError:  # pragma: no cover
        print("      (pdfplumber 未安装 → 跳过 PDF 文本抽取断言)")
        return
    with pdfplumber.open(str(e["out"] / "english_paper.pdf")) as pdf:
        text = "\n".join((page.extract_text() or "") for page in pdf.pages)
    assert text.strip(), "PDF 抽不出任何文本（可能是字体未正确嵌入）"
    assert "语法填空" in text, f"PDF 里应有中文分节名，实际前 200 字：{text[:200]}"
    assert "EN-APPROVED-G1" in text


def test_pdf_skipped_when_font_missing() -> None:
    """字体缺失分支：给警告、不生成 PDF、退出码仍为 0。"""
    e = e2e()
    data_dir = e["data_dir"]
    out_base = fresh_dir("nofont_out")
    saved = (ep.CJK_FONT_CANDIDATES, ep.CJK_BOLD_FONT_CANDIDATES)
    ep.CJK_FONT_CANDIDATES = [str(TMPROOT / "nope" / "msyh.ttc")]
    ep.CJK_BOLD_FONT_CANDIDATES = []
    try:
        assert ep.find_cjk_font() == (None, None)
        args = ep.build_parser().parse_args([
            "--subject", "english", "--data-dir", str(data_dir),
            "--outdir", str(out_base), "--date", E2E_DATE,
        ])
        buf = io.StringIO()
        with redirect_stdout(buf):
            code, summary = ep.run_export(args)
        assert code == 0, f"字体缺失应只告警，不应失败（退出码 {code}）"
        out = out_base / E2E_DATE
        assert (out / "english_paper.docx").is_file(), "字体缺失只影响 PDF"
        assert not list(out.glob("*.pdf")), "无中文字体时**不得**生成 PDF（避免方块字）"
        reasons = [s["reason"] for s in summary["skipped"] if s["item"] == "english_paper.pdf"]
        assert reasons and "字体" in reasons[0], f"必须有明确警告，实际 {summary['skipped']}"
        assert any("字体" in w for w in summary["warnings"])
    finally:
        ep.CJK_FONT_CANDIDATES, ep.CJK_BOLD_FONT_CANDIDATES = saved
        cleanup(out_base)


# --------------------------------------------------------------------- ⑥ Anki

def test_apkg_written_and_is_zip() -> None:
    e = e2e()
    path = e["out"] / "aca_llm.apkg"
    assert path.is_file(), "缺 aca_llm.apkg"
    assert path.stat().st_size > 0, "apkg 是空文件"
    assert zipfile.is_zipfile(str(path)), "apkg 应为 zip 容器"
    with zipfile.ZipFile(str(path)) as zf:
        names = set(zf.namelist())
        blob = zf.read("collection.anki2")
    assert "collection.anki2" in names, f"apkg 里应有 collection.anki2，实际 {names}"
    assert ep.ANKI_MODEL_ID and ep.ANKI_DECK_ID, "model/deck id 必须是固定常量"

    # guid = guid_for(question_uid) → 重复导出不会产生重复卡片（Anki 按 guid 去重）
    import genanki

    scratch = fresh_dir("apkg_guids") / "collection.anki2"
    scratch.write_bytes(blob)
    conn = sqlite3.connect(str(scratch))
    try:
        guids = {row[0] for row in conn.execute("SELECT guid FROM notes")}
    finally:
        conn.close()
    expected = {genanki.guid_for("aca_llm", uid) for uid in ("aca-s1", "aca-s2", "aca-m1")}
    assert guids == expected, f"apkg 的 note guid 应等于 guid_for(uid)（固定不变），实际 {guids}"


def test_apkg_skipped_without_choice_questions() -> None:
    d = fresh_dir("apkg_none")
    try:
        skipped: list[dict] = []
        ok = ep.write_apkg([q("x1", "writing", "只有主观题", subject="aca_llm")],
                           d / "aca_llm.apkg", {}, skipped)
        assert ok is False and not (d / "aca_llm.apkg").exists()
        assert skipped and "无 ACA 单选/多选" in skipped[0]["reason"]
    finally:
        cleanup(d)


# --------------------------------------------------------------------- ⑦ dry-run

def test_dry_run_writes_nothing() -> None:
    e = e2e()
    out_base = fresh_dir("dry_out")
    try:
        proc = run_cli("--data-dir", str(e["data_dir"]), "--outdir", str(out_base),
                       "--date", E2E_DATE, "--dry-run")
        assert proc.returncode == 0, f"dry-run 退出码 {proc.returncode}\n{proc.stderr[-800:]}"
        assert "[dry-run]" in proc.stdout
        summary = last_json(proc.stdout)
        assert summary["dry_run"] is True
        assert not (out_base / E2E_DATE).exists(), "dry-run 不得创建输出目录"
        assert list(out_base.rglob("*")) == [] or all(p.is_dir() for p in out_base.rglob("*")), \
            "dry-run 不得写任何文件"
        planned = {Path(f).name for f in summary["files"]}
        assert {"english_paper.docx", "all_questions.xlsx", "all_questions.csv", "aca_llm.apkg"} <= planned, \
            f"dry-run 应列出将生成的文件名，实际 {sorted(planned)}"
        assert summary["papers"]["english"]["actual_items"] == 6
    finally:
        cleanup(out_base)


# --------------------------------------------------------------------- 汇总 JSON / 其它

def test_cli_summary_shape_and_files_exist() -> None:
    e = e2e()
    summary = e["summary"]
    assert set(summary) >= {"papers", "files", "skipped"}
    assert summary["dry_run"] is False
    assert set(summary["papers"]) == {"english", "c_wuhan", "aca_llm"}, "c_hbpu 库缺失 → 不该出现在 papers"
    assert any(s.get("subject") == "c_hbpu" for s in summary["skipped"]), "缺库必须记入 skipped 并说明"
    assert len(summary["files"]) == len(set(summary["files"])), "文件列表不应有重复"
    for rel in summary["files"]:
        path = ROOT / rel
        assert path.is_file() and path.stat().st_size > 0, f"汇总里的文件不存在/为空：{rel}"


def test_c_subject_scores_marked_unknown() -> None:
    """C 语言两校官方配额未取得 → 分值写「未获取」，绝不编造。"""
    e = e2e()
    stats = e["summary"]["papers"]["c_wuhan"]
    assert stats["by_type"]["code_correction"]["per_item"] is None
    assert stats["by_type"]["programming"]["per_item"] is None
    from docx import Document
    text = "\n".join(p.text for p in Document(str(e["out"] / "c_wuhan_paper.docx")).paragraphs)
    assert "未获取" in text, "C 语言分值未知时文档里应写「未获取」"
    assert "CW-APPROVED-C1" in text


def test_unknown_question_type_is_not_silently_dropped() -> None:
    """taxonomy 之外的题型（如英语里混进的选择题）**默认不编入卷面**，但必须**明确标注**数量与题型名，
    且可以通过 `include_unknown=True` 显式编入 —— **任何情况下都不得静默丢题**。

    ⚠️ 契约变更（真题导入后暴露的问题）：旧实现把越界题型无条件追加成分节且**没有配额上限**，
    结果英语卷面从 43 题暴涨到 431 题（完形填空/语音题/情景对话等新题型涌入）。
    现在默认不编入正文，卷末给一条 note 说明；要编入必须显式开关。
    """
    d = fresh_dir("unknown_type")
    try:
        db = make_db(d / "english.sqlite", [
            q("ok1", "grammar_fill", "EN-OK-1 正常题", subject="english"),
            q("bad1", "single_choice", "EN-OUT-OF-SCOPE-1 越界选择题", subject="english",
              options={"A": "甲", "B": "乙"}, answer="A"),
        ])
        ata = ep.load_taxonomy()
        warns: list[str] = []

        # (1) 默认：越界题型**不进卷**，但卷末有 note + 警告里有题型名
        built = ep.build_document("english", ep.subject_meta(ata, "english"),
                                  ep.question_types(ata, "english"),
                                  ep.fetch_questions(db, ["approved"]),
                                  statuses=["approved"], generated_at="2026-01-01 00:00", warnings=warns)
        assert len(built["selected"]) == 1, f"默认只应收录大纲内题型，实际 {len(built['selected'])} 题"
        assert built["stats"]["uncategorized_items"] == 1, "越界题量仍须照实统计"
        texts = [b["text"] for b in built["blocks"]]
        assert not any("EN-OUT-OF-SCOPE-1" in t for t in texts), "默认不应把越界题渲染进卷面"
        assert any("大纲之外" in t and "single_choice" in t for t in texts), \
            "必须在卷末明确标注越界题型的数量与名称（不得静默丢弃）"
        assert warns and "single_choice" in warns[0], "越界题型必须给警告"

        # (2) 显式开关：编入正文并渲染题干与选项
        warns2: list[str] = []
        built2 = ep.build_document("english", ep.subject_meta(ata, "english"),
                                   ep.question_types(ata, "english"),
                                   ep.fetch_questions(db, ["approved"]),
                                   statuses=["approved"], generated_at="2026-01-01 00:00",
                                   warnings=warns2, include_unknown=True)
        assert len(built2["selected"]) == 2, "include_unknown=True 时应把越界题编入"
        t2 = [b["text"] for b in built2["blocks"]]
        assert any("single_choice" in t for t in t2), "越界题型应有自己的分节标题"
        assert any("EN-OUT-OF-SCOPE-1" in t for t in t2), "越界题型的题干必须被渲染"
        assert any("A. 甲" == t for t in t2), "越界题的选项也必须被渲染"
    finally:
        cleanup(d)


def test_compile() -> None:
    import py_compile

    py_compile.compile(str(ROOT / "scripts" / "export_papers.py"), doraise=True)


if __name__ == "__main__":
    fns = [v for k, v in sorted(globals().items()) if k.startswith("test_") and callable(v)]
    failed = 0
    for fn in fns:
        try:
            fn()
            print(f"  PASS {fn.__name__}")
        except AssertionError as exc:
            failed += 1
            print(f"  FAIL {fn.__name__}: {exc}")
        except Exception as exc:  # noqa: BLE001
            failed += 1
            print(f"  ERROR {fn.__name__}: {type(exc).__name__}: {exc}")
    shutil.rmtree(TMPROOT, ignore_errors=True)
    print(f"{len(fns) - failed}/{len(fns)} passed")
    raise SystemExit(1 if failed else 0)
